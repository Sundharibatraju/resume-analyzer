"""
Extracts raw text from uploaded PDF and DOCX resume files.

PDF: pdfplumber is tried first (better layout/whitespace handling),
falling back to PyPDF2 if pdfplumber fails on a given file (e.g. some
malformed or image-heavy PDFs raise errors in one library but not the other).

DOCX: python-docx, pulling paragraph text and table cell text.
"""
import pdfplumber
import PyPDF2
import docx
from app.utils.logger import get_logger

logger = get_logger(__name__)


class DocumentExtractionError(Exception):
    pass


def extract_text_from_pdf(file_path: str) -> str:
    text_parts = []

    # Attempt 1: pdfplumber
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        text = "\n".join(text_parts).strip()
        if text:
            return text
    except Exception as e:
        logger.warning(f"pdfplumber failed for {file_path}: {e}")

    # Attempt 2 (fallback): PyPDF2
    try:
        text_parts = []
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        text = "\n".join(text_parts).strip()
        if text:
            return text
    except Exception as e:
        logger.error(f"PyPDF2 fallback also failed for {file_path}: {e}")

    raise DocumentExtractionError(
        "Could not extract text from this PDF. It may be scanned/image-based or corrupted."
    )


def extract_text_from_docx(file_path: str) -> str:
    try:
        document = docx.Document(file_path)
        text_parts = [p.text for p in document.paragraphs if p.text.strip()]

        # Also pull text from tables (some resumes use table layouts)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        text = "\n".join(text_parts).strip()
        if not text:
            raise DocumentExtractionError("DOCX file contains no extractable text.")
        return text
    except DocumentExtractionError:
        raise
    except Exception as e:
        logger.error(f"python-docx failed for {file_path}: {e}")
        raise DocumentExtractionError(f"Could not extract text from this DOCX file: {e}")


def extract_text(file_path: str, file_type: str) -> str:
    file_type = file_type.lower()
    if file_type == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_type == "docx":
        return extract_text_from_docx(file_path)
    else:
        raise DocumentExtractionError(f"Unsupported file type: {file_type}")
